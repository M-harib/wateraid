"""Async ingest of pack docs into Units, so both tools read one corpus."""
from __future__ import annotations

import asyncio

from ..schema import DocRecord, SourceRef, Unit


def _parse_pdf(doc_id: str, version: str | None, path: str) -> list[Unit]:
    from pypdf import PdfReader

    reader = PdfReader(path)
    units: list[Unit] = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if not text:
            continue
        ref = SourceRef(doc_id=doc_id, version=version, page=i)
        units.append(Unit(doc_id=doc_id, index=i, ref=ref, text=text))
    return units


def _parse_xlsx(doc_id: str, version: str | None, path: str) -> list[Unit]:
    import openpyxl
    from openpyxl.utils import get_column_letter

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    units: list[Unit] = []
    idx = 0
    try:
        for ws in wb.worksheets:
            header: list[str] | None = None
            for r, row in enumerate(ws.iter_rows(), start=1):
                values = [("" if c.value is None else str(c.value)).strip() for c in row]
                if not any(values):
                    continue
                nonempty = [v for v in values if v]
                # first multi-column row is the header; merge it into later rows
                if header is None and len(nonempty) > 1:
                    header = values
                last_col = max((j for j, v in enumerate(values) if v), default=0)
                cells = f"A{r}:{get_column_letter(last_col + 1)}{r}"
                if header and header != values:
                    pairs = [f"{h}: {v}" for h, v in zip(header, values) if v]
                    text = " | ".join(pairs) if pairs else " | ".join(nonempty)
                else:
                    text = " | ".join(nonempty)
                idx += 1
                ref = SourceRef(doc_id=doc_id, version=version, sheet=ws.title, cells=cells)
                units.append(Unit(doc_id=doc_id, index=idx, ref=ref, text=text))
    finally:
        wb.close()
    return units


async def ingest_doc(rec: DocRecord) -> list[Unit]:
    if not (rec.present and rec.in_corpus and rec.path):
        return []
    # ponytail: thread pool suits IO-bound parsing; ProcessPoolExecutor if CPU-bound.
    if rec.kind == "pdf":
        return await asyncio.to_thread(_parse_pdf, rec.doc_id, rec.version, rec.path)
    if rec.kind == "excel":
        return await asyncio.to_thread(_parse_xlsx, rec.doc_id, rec.version, rec.path)
    return []


async def ingest_all(records) -> list[Unit]:
    """Ingest every in-corpus doc concurrently; flatten to one Unit list."""
    batches = await asyncio.gather(*(ingest_doc(r) for r in records))
    return [u for batch in batches for u in batch]
